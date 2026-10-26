from __future__ import annotations

import shutil
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.helpers import cards as cards_module
from src.helpers.cards import (
    build_cards,
    card_filename,
    render_docx_bytes,
    write_cards_zip,
)
from src.helpers.data_models import InnerDict, NameKey, OuterDict
from src.helpers.vars import (
    DRAW_LABEL,
    KTP_FILENAME_COL,
    KTP_FIRST_NAME_ORIG_COLNAME_COL,
    KTP_HCR_FILENAME_COL,
    KTP_HCR_FILENAME_COL_LEGACY,
    KTP_LAST_NAME_ORIG_COLNAME_COL,
    KTP_SSN_TOP_OLDEST_PAPERS_COL,
    KTP_SSNP_PAPERID_URL_COL,
    OPENALEX_TITLE_COL,
    SSNP_DATE_COL,
)

UNDERSCORE_FIELD_LABEL = "ktp.field_with_underscores_"
RENDERED_UNDERSCORE_FIELD_LABEL = f"**`{UNDERSCORE_FIELD_LABEL}`**"
UNDERSCORE_FILENAME = "source_file.xlsx"
RENDERED_UNDERSCORE_FILENAME = f"#### {KTP_FILENAME_COL}: `{UNDERSCORE_FILENAME}`"
ROUNDTRIP_CARD_NAME = "1_Ada_Lovelace"


class DummyProcedure:
    dataset_id_field = "ktp.source_key"


def test_card_filename_matches_card_archive_member_stem() -> None:
    assert (
        card_filename(
            draw_label="1, pilot.2",
            first_name="Ada",
            last_name="Lovelace-Smith",
        )
        == "1_pilot2_Ada_LovelaceSmith"
    )


def test_render_docx_bytes_uses_exact_markdown_and_cleans_temporary_files(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    markdown = "## Exact card\n\nLine with _underscores_ and UTF-8: naïve\n"
    expected_docx = b"PK\x03\x04canonical-docx"
    reference_docx = tmp_path / "reference.docx"
    reference_docx.write_bytes(b"reference")
    temporary_paths: tuple[Path, Path] | None = None

    def render_docx(
        md_path: Path,
        docx_path: Path,
        supplied_reference_docx: Path,
    ) -> Path:
        nonlocal temporary_paths
        temporary_paths = (md_path, docx_path)
        assert md_path.read_bytes() == markdown.encode("utf-8")
        assert supplied_reference_docx == reference_docx
        docx_path.write_bytes(expected_docx)
        return docx_path

    monkeypatch.setattr(cards_module, "_render_docx", render_docx)

    assert render_docx_bytes(markdown, reference_docx) == expected_docx
    assert temporary_paths is not None
    assert all(not path.exists() for path in temporary_paths)


def test_build_cards_includes_intro_and_fun_fact() -> None:
    name_key = NameKey(first_name="Ada", last_name="Lovelace")
    outer = OuterDict.from_name_keys([name_key])

    data = {
        DRAW_LABEL: 1,
        KTP_FILENAME_COL: "2019_HCR.xlsx",
        KTP_FIRST_NAME_ORIG_COLNAME_COL: "First Name",
        KTP_LAST_NAME_ORIG_COLNAME_COL: "Last Name",
        KTP_SSN_TOP_OLDEST_PAPERS_COL: (
            f'[{{"{SSNP_DATE_COL}":"1903-01-01",'
            f'"{OPENALEX_TITLE_COL}":"Early work",'
            f'"{KTP_SSNP_PAPERID_URL_COL}":"https://openalex.org/W1568216332"}}]'
        ),
        "note": "hello",
        "excluded": "skip",
    }
    outer.add_inner(name_key, InnerDict.from_mapping(data, DummyProcedure()))

    cards = build_cards(
        outer,
        total_draws=10,
        intro="## Introduction\nDate of report: 2026-02-02",
        excluded_cols={"excluded"},
    )

    assert "1_Ada_Lovelace" in cards
    card = cards["1_Ada_Lovelace"]
    assert "## Introduction" in card
    assert "Fun fact" in card
    assert "excluded" not in card
    assert "Draw #1 of 10" in card
    assert (
        f"**`{KTP_SSN_TOP_OLDEST_PAPERS_COL}`**: "
        f"`{data[KTP_SSN_TOP_OLDEST_PAPERS_COL]}`"
    ) in card
    assert "Early work" in card
    assert "https://openalex.org/W1568216332" in card


def test_build_cards_field_renderer_can_replace_missing_values_and_omit_fields() -> None:
    name_key = NameKey(first_name="Ada", last_name="Lovelace")
    inner = InnerDict.from_mapping(
        {
            DRAW_LABEL: 1,
            "table_position": None,
            "standalone_value": "must be moved into the table",
            "other_value": "still visible",
            "excluded": "must stay excluded",
        },
        DummyProcedure(),
    )
    outer = OuterDict(data={name_key.to_json_key(): [inner]})
    seen: list[str] = []

    def render_field(candidate: InnerDict, column: str) -> str | None:
        assert candidate is inner
        seen.append(column)
        if column == "table_position":
            return "| field | value |\n| --- | --- |\n| row | replacement |\n\n"
        if column == "standalone_value":
            return ""
        return None

    card = next(iter(build_cards(
        outer,
        total_draws=1,
        intro="Introduction",
        excluded_cols={"excluded"},
        field_renderer=render_field,
    ).values()))

    assert "| row | replacement |" in card
    assert "standalone_value" not in card
    assert "must be moved into the table" not in card
    assert "**`other_value`**: still visible" in card
    assert "excluded" not in card
    assert "excluded" not in seen
    assert "table_position" in seen


def test_render_docx_bytes_preserves_table_linebreaks_and_adds_plain_borders(
    tmp_path: Path,
) -> None:
    if shutil.which("pandoc") is None:
        pytest.skip("pandoc not available")

    reference_docx = tmp_path / "reference.docx"
    Document().save(str(reference_docx))
    rendered = render_docx_bytes(
        "| field | value |\n| --- | --- |\n| row | first<br>second |\n",
        reference_docx,
    )
    document = Document(BytesIO(rendered))
    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(1, 1).text == "first\nsecond"
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        assert border is not None
        assert border.get(qn("w:val")) == "single"
        assert border.get(qn("w:sz")) == "4"
        assert border.get(qn("w:color")) == "000000"


def test_plain_table_borders_leave_existing_borders_and_other_docx_parts_unchanged() -> None:
    document = Document()
    existing = document.add_table(rows=1, cols=1)
    borders = OxmlElement("w:tblBorders")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "double")
    top.set(qn("w:color"), "FF0000")
    borders.append(top)
    existing._tbl.tblPr.append(borders)
    document.add_table(rows=1, cols=1)
    source = BytesIO()
    document.save(source)

    revised = cards_module._add_plain_table_borders(source.getvalue())
    assert cards_module._add_plain_table_borders(revised) == revised
    with ZipFile(BytesIO(source.getvalue())) as before, ZipFile(BytesIO(revised)) as after:
        assert before.namelist() == after.namelist()
        for part in before.namelist():
            if part != "word/document.xml":
                assert before.read(part) == after.read(part)

    tables = Document(BytesIO(revised)).tables
    preserved = tables[0]._tbl.tblPr.find(qn("w:tblBorders"))
    assert preserved is not None
    assert len(preserved) == 1
    assert preserved[0].get(qn("w:val")) == "double"
    assert preserved[0].get(qn("w:color")) == "FF0000"
    assert tables[1]._tbl.tblPr.find(qn("w:tblBorders")) is not None


def test_write_cards_zip_txt(tmp_path: Path) -> None:
    cards = {"Ada_Lovelace": "card content", "Grace_Hopper": "more"}
    zip_path = write_cards_zip(
        cards,
        tmp_path,
        "cards.zip",
        output_format="txt",
        reference_docx=tmp_path / "ref.docx",
    )

    assert zip_path.exists()
    with ZipFile(zip_path, "r") as zipf:
        assert sorted(zipf.namelist()) == ["Ada_Lovelace.txt", "Grace_Hopper.txt"]


def test_write_cards_zip_docx_skips_without_pandoc(tmp_path: Path) -> None:
    if shutil.which("pandoc") is None:
        pytest.skip("pandoc not available")

    cards = {"Ada_Lovelace": "card content"}
    ref_docx = Path("resources/pandoc-custom-reference.docx")
    if not ref_docx.exists():
        pytest.skip("reference docx not available")

    zip_path = write_cards_zip(
        cards,
        tmp_path,
        "cards_docx.zip",
        output_format="docx",
        reference_docx=ref_docx,
    )

    assert zip_path.exists()
    with ZipFile(zip_path, "r") as zipf:
        assert zipf.namelist() == ["Ada_Lovelace.docx"]


def test_underscore_field_labels_round_trip_in_txt_and_docx(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    name_key = NameKey(first_name="Ada", last_name="Lovelace")
    outer = OuterDict.from_name_keys([name_key])
    outer.add_inner(
        name_key,
        InnerDict.from_mapping(
            {
                DRAW_LABEL: 1,
                KTP_FILENAME_COL: UNDERSCORE_FILENAME,
                KTP_HCR_FILENAME_COL: UNDERSCORE_FILENAME,
                KTP_HCR_FILENAME_COL_LEGACY: UNDERSCORE_FILENAME,
                UNDERSCORE_FIELD_LABEL: "literal value",
            },
            DummyProcedure(),
        ),
    )
    cards = build_cards(
        outer,
        total_draws=1,
        intro="## Introduction",
        excluded_cols=set(),
    )
    card = cards[ROUNDTRIP_CARD_NAME]
    assert RENDERED_UNDERSCORE_FIELD_LABEL in card
    assert RENDERED_UNDERSCORE_FILENAME in card
    assert f"**{KTP_FILENAME_COL}**: `{UNDERSCORE_FILENAME}`" in card
    assert f"**`{KTP_HCR_FILENAME_COL}`**: `{UNDERSCORE_FILENAME}`" in card
    assert f"**{KTP_HCR_FILENAME_COL_LEGACY}**: {UNDERSCORE_FILENAME}" in card

    txt_zip = write_cards_zip(
        cards,
        tmp_path,
        "cards_txt.zip",
        output_format="txt",
        reference_docx=tmp_path / "unused-reference.docx",
    )
    with ZipFile(txt_zip) as archive:
        txt_card = archive.read(f"{ROUNDTRIP_CARD_NAME}.txt").decode("utf-8")
    assert RENDERED_UNDERSCORE_FIELD_LABEL in txt_card
    assert RENDERED_UNDERSCORE_FILENAME in txt_card

    captured_markdown: list[str] = []

    def render_docx(
        md_path: Path,
        docx_path: Path,
        _reference_docx: Path,
    ) -> Path:
        captured_markdown.append(md_path.read_text(encoding="utf-8"))
        docx_path.write_bytes(b"docx")
        return docx_path

    reference_docx = tmp_path / "reference.docx"
    reference_docx.write_bytes(b"reference")
    monkeypatch.setattr(cards_module, "_render_docx", render_docx)
    docx_zip = write_cards_zip(
        cards,
        tmp_path,
        "cards_docx.zip",
        output_format="docx",
        reference_docx=reference_docx,
        docx_workers=1,
    )
    assert captured_markdown == [card]
    assert RENDERED_UNDERSCORE_FIELD_LABEL in captured_markdown[0]
    assert RENDERED_UNDERSCORE_FILENAME in captured_markdown[0]
    with ZipFile(docx_zip) as archive:
        assert archive.namelist() == [f"{ROUNDTRIP_CARD_NAME}.docx"]
